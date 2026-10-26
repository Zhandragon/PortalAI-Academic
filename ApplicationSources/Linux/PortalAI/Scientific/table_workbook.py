"""Bounded value-only Excel ingestion; no macro/formula execution."""
import io,sys,zipfile
from Scientific.extensions_common import read,integer,keys,sha

def frame(p,root):
    from Scientific.biology_tool import Context
    context=Context(root);data=read(context,p['path'],32*1024**2)
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        items=z.infolist()
        if len(items)>10000 or sum(x.file_size for x in items)>128*1024**2:raise ValueError('Workbook expanded size exceeds limits')
        for item in items:
            if item.file_size>32*1024**2:raise ValueError('Workbook member exceeds 32 MiB')
            if item.filename.endswith('.xml'):
                content=z.read(item)
                if b'<!ENTITY' in content.upper() or b'<!DOCTYPE' in content.upper():raise ValueError('Workbook XML entities/DTDs refused')
    packages=context.root/'Runtime/ScienceExtras/packages'
    if str(packages) not in sys.path:sys.path.insert(0,str(packages))
    import openpyxl,pandas as pd
    formula_book=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
    values_book=None
    try:
        sheet=p.get('sheet',0)
        if isinstance(sheet,str):ws=formula_book[sheet]
        else:ws=formula_book.worksheets[integer(sheet,'sheet',0,len(formula_book.worksheets)-1)]
        limit=integer(p.get('max_rows',100000),'max_rows',1,100000)
        if ws.max_column>200 or ws.max_row>limit+1 or ws.max_column*ws.max_row>2000000:raise ValueError('Worksheet exceeds row/column/cell limit; select a smaller sheet')
        rows=list(ws.iter_rows());formulas=sum(c.data_type=='f' for row in rows for c in row)
        if formulas:
            if p.get('use_cached_formulas') is not True:raise ValueError('Workbook contains formulas. Recalculate/save in Excel, then explicitly set use_cached_formulas:true to read cached values; formulas are never executed here.')
            values_book=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=True,keep_links=False);cached=list(values_book[ws.title].iter_rows())
            for r,row in enumerate(rows):
                for c,cell in enumerate(row):
                    if cell.data_type=='f' and cached[r][c].value is None:raise ValueError('Formula has no cached value; recalculate/save the workbook first')
            rows=cached
        values=[[c.value for c in row] for row in rows]
        if not values:raise ValueError('Empty worksheet')
        header=[str(v) if v is not None else '' for v in values[0]]
        if any(not v for v in header) or len(set(header))!=len(header):raise ValueError('First row requires unique nonempty column names')
        result=pd.DataFrame(values[1:],columns=header)
        result.attrs.update(sheet=ws.title,workbook_sha256=sha(data),formula_cells=formulas,formula_note='Cached values may be stale; no calculation or external link refresh was performed.')
        context.check();return result
    finally:
        formula_book.close()
        if values_book:values_book.close()

def run(p,context):
    keys(p,{'path','sheet','max_rows','use_cached_formulas','preview_rows'})
    df=frame(p,context.root);n=integer(p.get('preview_rows',20),'preview_rows',1,100)
    return {'rows':len(df),'columns':list(df.columns),'preview':__import__('json').loads(df.head(n).to_json(orient='records',date_format='iso')),**df.attrs},{'table.csv':df.to_csv(index=False)}
